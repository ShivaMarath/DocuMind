"""Document loading, chunking and idempotent upserts into Qdrant."""
from __future__ import annotations

import hashlib
import uuid
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import get_settings

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    """Load a file into Documents (one per PDF page, otherwise one per file)."""
    path = Path(path)
    ext = path.suffix.lower()
    base_meta = {"source": str(path), "source_name": path.name}

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [
            Document(page_content=page.extract_text() or "", metadata={**base_meta, "page": i + 1})
            for i, page in enumerate(reader.pages)
            if (page.extract_text() or "").strip()
        ]
    if ext in {".txt", ".md"}:
        return [Document(page_content=path.read_text(encoding="utf-8", errors="ignore"), metadata=base_meta)]
    if ext == ".docx":
        import docx

        text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        return [Document(page_content=text, metadata=base_meta)]
    raise ValueError(f"Unsupported file type: {ext} (supported: {sorted(SUPPORTED_EXTENSIONS)})")


def split_documents(docs: list[Document], chunk_size: int | None = None, chunk_overlap: int | None = None) -> list[Document]:
    s = get_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size or s.chunk_size,
        chunk_overlap=chunk_overlap if chunk_overlap is not None else s.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(docs)
    counters: dict[tuple, int] = {}
    for c in chunks:
        key = (c.metadata.get("source_name"), c.metadata.get("page"))
        counters[key] = counters.get(key, 0) + 1
        c.metadata["chunk_index"] = counters[key] - 1
    return [c for c in chunks if c.page_content.strip()]


def chunk_id(doc: Document) -> str:
    """Deterministic ID: re-ingesting identical content overwrites instead of duplicating."""
    digest = hashlib.sha1(doc.page_content.encode("utf-8")).hexdigest()
    m = doc.metadata
    raw = f"{m.get('source_name')}|{m.get('page')}|{m.get('chunk_index')}|{digest}"
    return str(uuid.uuid5(uuid.NAMESPACE_URL, raw))


def collect_files(paths: list[str | Path]) -> list[Path]:
    files: list[Path] = []
    for p in map(Path, paths):
        if p.is_dir():
            files += sorted(f for f in p.rglob("*") if f.suffix.lower() in SUPPORTED_EXTENSIONS)
        elif p.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(p)
        else:
            print(f"  skipping unsupported file: {p}")
    return files


def ingest_paths(paths: list[str | Path], vectorstore=None) -> dict:
    """Index files. One bad file never aborts the batch; problems are reported per file."""
    from .vectorstore import delete_source, get_vectorstore

    vs = vectorstore or get_vectorstore()
    summary = {"files": 0, "chunks": 0, "details": []}
    for f in collect_files(paths):
        try:
            chunks = split_documents(load_file(f))
            if not chunks:
                summary["details"].append({
                    "file": f.name, "chunks": 0,
                    "note": "no extractable text (scanned/image-only PDF? OCR is not supported)"})
                continue
            ids = [chunk_id(c) for c in chunks]
            vs.add_documents(chunks, ids=ids)          # write new chunks first...
            delete_source(f.name, keep_ids=ids)        # ...then drop stale chunks from older versions
            summary["files"] += 1
            summary["chunks"] += len(chunks)
            summary["details"].append({"file": f.name, "chunks": len(chunks)})
        except Exception as e:  # noqa: BLE001 - surfaced to the caller/UI
            summary["details"].append({"file": f.name, "chunks": 0, "error": f"{type(e).__name__}: {e}"})
    return summary
